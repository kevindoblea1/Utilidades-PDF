#!/usr/bin/env python3
import argparse, os, zipfile, sys, io, math
import pandas as pd

# Opcionales: si faltan, el script sigue con lo disponible
try:
    from docx import Document
    HAS_PYDOCX = True
except Exception:
    HAS_PYDOCX = False

try:
    from PIL import Image as PILImage
    HAS_PIL = True
except Exception:
    HAS_PIL = False

def choose_engine():
    try:
        import xlsxwriter  # noqa
        return "xlsxwriter"
    except Exception:
        return "openpyxl"

def read_tables_docx(path):
    """DataFrames desde tablas estándar de Word."""
    dfs = []
    if not HAS_PYDOCX:
        return dfs
    doc = Document(path)
    for t in doc.tables:
        rows = []
        max_cols = max((len(r.cells) for r in t.rows), default=0)
        for r in t.rows:
            rows.append([c.text.strip() for c in r.cells] + [""] * (max_cols - len(r.cells)))
        if not rows:
            continue
        header = [h.strip() for h in rows[0]]
        if sum(bool(x) for x in header) >= max(1, int(len(header) * 0.5)):
            df = pd.DataFrame(rows[1:], columns=header)
        else:
            df = pd.DataFrame(rows)
        df.replace("", pd.NA, inplace=True)
        df = df.ffill().bfill()
        dfs.append(df)
    return dfs

def read_paragraphs_xml(path):
    """Párrafos desde word/document.xml (incluye cuadros de texto)."""
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(path) as z:
        xml = z.read('word/document.xml')
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    root = ET.fromstring(xml)
    lines = []
    for p in root.findall('.//w:p', ns):
        texts = [t.text for t in p.findall('.//w:t', ns) if t.text]
        line = "".join(texts).strip()
        if line:
            lines.append(line)
    # Dedup suave
    cleaned, prev = [], None
    for ln in lines:
        if ln != prev:
            cleaned.append(ln)
        prev = ln
    return cleaned

def extract_images_docx(path):
    """
    Devuelve lista de dicts:
    {'name': 'img1.png', 'bytes': b'..', 'fmt': 'PNG', 'w': px, 'h': px}
    """
    out = []
    with zipfile.ZipFile(path) as z:
        members = [m for m in z.namelist() if m.startswith('word/media/')]
        for m in sorted(members):
            data = z.read(m)
            name = os.path.basename(m)
            ext = os.path.splitext(name)[1].lower()
            fmt = ext.replace('.', '').upper()

            # Normaliza a PNG/JPEG; convierte si hace falta
            if HAS_PIL:
                try:
                    with PILImage.open(io.BytesIO(data)) as im:
                        w, h = im.size
                        if fmt not in ('PNG', 'JPG', 'JPEG'):
                            bio = io.BytesIO()
                            im.save(bio, format='PNG')
                            data = bio.getvalue()
                            fmt = 'PNG'
                            w, h = PILImage.open(io.BytesIO(data)).size
                        out.append({'name': name, 'bytes': data, 'fmt': fmt, 'w': w, 'h': h})
                        continue
                except Exception:
                    pass  # cae al fallback
            # Fallback sin Pillow: deja tal cual (xlsxwriter soporta JPG/PNG)
            out.append({'name': name, 'bytes': data, 'fmt': fmt, 'w': None, 'h': None})
    return out

def write_xlsx(dfs, paragraphs, images, out_path, include_combined=True):
    engine = choose_engine()
    with pd.ExcelWriter(out_path, engine=engine) as w:
        wrote_tables = False

        # 1) Tablas
        for i, df in enumerate(dfs, start=1):
            sheet = f"Tabla{i}"
            df.to_excel(w, sheet_name=sheet, index=False)
            wrote_tables = True

        # 2) Combinado
        if wrote_tables and include_combined and len(dfs) > 1:
            combo = pd.concat(dfs, ignore_index=True)
            combo.to_excel(w, sheet_name="Combinado", index=False)

        # 3) Texto
        if paragraphs:
            pd.DataFrame({"Texto": paragraphs}).to_excel(w, sheet_name="Texto", index=False)
        elif not wrote_tables and not images:
            pd.DataFrame({"Mensaje": ["El documento no contiene tablas, texto ni imágenes."]}).to_excel(
                w, sheet_name="Mensaje", index=False
            )

        # 4) Imágenes
        if images:
            if engine == "xlsxwriter":
                # Crear hoja manualmente
                ws = w.book.add_worksheet("Imagenes")
                w.sheets["Imagenes"] = ws
                row = 0
                max_width_px = 900  # ancho deseado
                for idx, img in enumerate(images, start=1):
                    ws.write(row, 0, f"{idx}. {img['name']}")
                    # Escala
                    x_scale = y_scale = 1.0
                    if img.get('w') and img.get('h'):
                        if img['w'] > max_width_px:
                            x_scale = max_width_px / float(img['w'])
                            y_scale = x_scale
                        est_rows = int(math.ceil((img['h'] * y_scale) / 18.0)) + 2  # 18px aprox altura
                    else:
                        est_rows = 30
                    bio = io.BytesIO(img['bytes'])
                    ws.insert_image(row + 1, 0, img['name'], {'image_data': bio, 'x_scale': x_scale, 'y_scale': y_scale})
                    row += est_rows
            else:
                # openpyxl
                ws = w.book.create_sheet("Imagenes")
                w.sheets["Imagenes"] = ws
                try:
                    from openpyxl.drawing.image import Image as XLImage
                    row = 1
                    max_width_px = 900
                    for idx, img in enumerate(images, start=1):
                        ws.cell(row=row, column=1, value=f"{idx}. {img['name']}")
                        # openpyxl requiere PIL para métricas correctas
                        if HAS_PIL:
                            with PILImage.open(io.BytesIO(img['bytes'])) as im:
                                w0, h0 = im.size
                                scale = 1.0
                                if w0 > max_width_px:
                                    scale = max_width_px / float(w0)
                                    im = im.resize((int(w0*scale), int(h0*scale)))
                                bio = io.BytesIO()
                                im.save(bio, format='PNG')
                                bio.seek(0)
                                xlimg = XLImage(bio)
                        else:
                            # Si no hay PIL, intenta insertar tal cual (PNG/JPEG)
                            bio = io.BytesIO(img['bytes'])
                            xlimg = XLImage(bio)
                        ws.add_image(xlimg, f"A{row+1}")
                        # estimar salto de filas
                        est_rows = 30
                        if HAS_PIL:
                            try:
                                hpx = xlimg.height
                                est_rows = int(math.ceil(hpx / 18.0)) + 2
                            except Exception:
                                pass
                        row += est_rows
                except Exception:
                    # Si no se pudieron convertir, al menos deja listado
                    pd.DataFrame([i['name'] for i in images], columns=["Imagen"]).to_excel(
                        w, sheet_name="Imagenes", index=False
                    )

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input", help=".docx de entrada")
    ap.add_argument("-o", "--output", help="Ruta de salida .xlsx")
    ap.add_argument("--no-combined", action="store_true", help="No crear hoja 'Combinado'")
    args = ap.parse_args()

    in_path = args.input
    out_path = args.output or os.path.splitext(in_path)[0] + ".xlsx"

    dfs = read_tables_docx(in_path)
    paragraphs = read_paragraphs_xml(in_path)  # extrae texto aunque no haya tablas
    images = extract_images_docx(in_path)

    write_xlsx(dfs, paragraphs, images, out_path, include_combined=(not args.no_combined))
    print(out_path)

if __name__ == "__main__":
    sys.exit(main())
